# -*- coding: utf-8 -*-
"""星火三部曲 · 项目介绍演示文稿生成器（python-pptx）。"""
import math
import os
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECK = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "星火三部曲·项目介绍.pptx"

BG = RGBColor(0x07, 0x0B, 0x18)
PANEL = RGBColor(0x12, 0x24, 0x49)
GOLD = RGBColor(0xF6, 0xDC, 0x9B)
GOLD2 = RGBColor(0xE6, 0xAD, 0x45)
GOLD3 = RGBColor(0xC9, 0x8A, 0x2B)
FIRE = RGBColor(0xF2, 0x60, 0x1D)
FIRE2 = RGBColor(0xFF, 0x9A, 0x5A)
INK = RGBColor(0xEE, 0xF3, 0xFF)
MUTED = RGBColor(0x9F, 0xB2, 0xD8)
LINE = RGBColor(0x2C, 0x3E, 0x68)

SERIF = "Noto Serif SC"
SANS = "Microsoft YaHei"

W, H = 13.333, 7.5
LEFT = 0.95
CONTENT_W = 11.45


# ---------------------------------------------------------------- background
def make_background(path: str, cover: bool) -> None:
    w, h = 2560, 1440
    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32)

    def glow(cx, cy, radius, color, strength):
        d = np.sqrt(((xx - cx) / radius) ** 2 + ((yy - cy) / radius) ** 2)
        return (np.clip(1 - d, 0, 1) ** 2 * strength)[..., None]

    base = np.zeros((h, w, 3), np.float32)
    base[:] = (7, 11, 24)
    layers = [
        (w * 0.86, h * -0.08, w * 0.62, (230, 173, 69), 0.30 if cover else 0.18),
        (w * -0.08, h * 0.82, w * 0.58, (242, 96, 29), 0.20 if cover else 0.10),
        (w * 0.42, h * 0.30, w * 0.50, (43, 69, 127), 0.16),
    ]
    for cx, cy, r, c, s in layers:
        a = glow(cx, cy, r, c, s)
        base = base * (1 - a) + np.array(c, float) * a

    vig = np.clip(1.25 - (((xx / w - 0.5) ** 2 + (yy / h - 0.5) ** 2) ** 0.5) * 1.5, 0.35, 1)
    base *= vig[..., None]

    img = Image.fromarray(np.clip(base, 0, 255).astype(np.uint8), "RGB").convert("RGBA")
    stars = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(stars)
    rng = np.random.default_rng(20261022 if cover else 20261023)
    for _ in range(300 if cover else 190):
        x, y = rng.uniform(0, w), rng.uniform(0, h)
        r = rng.uniform(0.7, 2.4) * (1.5 if cover else 1.0)
        a = int(rng.uniform(30, 190) * (1.0 if cover else 0.7))
        tint = (255, 246, 220) if rng.random() < 0.28 else (222, 232, 255)
        d.ellipse([x - r, y - r, x + r, y + r], fill=tint + (a,))
        if r > 1.9 and rng.random() < 0.5:
            d.line([x - r * 4, y, x + r * 4, y], fill=tint + (int(a * 0.45),), width=1)
            d.line([x, y - r * 4, x, y + r * 4], fill=tint + (int(a * 0.45),), width=1)
    img = Image.alpha_composite(img, stars)
    img.convert("RGB").save(path, quality=92)
    print("bg:", path)


def rounded_shot(src: Path, dst: Path, aspect=(16, 10), radius=34) -> None:
    im = Image.open(src).convert("RGB")
    tw, th = aspect
    target = tw / th
    w, h = im.size
    if w / h > target:
        nw = int(h * target)
        im = im.crop(((w - nw) // 2, 0, (w + nw) // 2, h))
    else:
        nh = int(w / target)
        im = im.crop((0, (h - nh) // 2, w, (h + nh) // 2))
    im = im.resize((1800, int(1800 * th / tw)), Image.LANCZOS).convert("RGBA")

    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, im.size[0] - 1, im.size[1] - 1], radius, fill=255)
    im.putalpha(mask)

    pad = 90
    canvas = Image.new("RGBA", (im.size[0] + pad * 2, im.size[1] + pad * 2), (0, 0, 0, 0))
    shadow = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    sh = Image.new("RGBA", im.size, (0, 0, 0, 190))
    sh.putalpha(mask.point(lambda v: int(v * 0.55)))
    shadow.paste(sh, (pad, pad + 22), sh)
    shadow = shadow.filter(ImageFilter.GaussianBlur(26))
    canvas = Image.alpha_composite(canvas, shadow)
    canvas.paste(im, (pad, pad), im)

    border = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(border).rounded_rectangle(
        [pad, pad, pad + im.size[0] - 1, pad + im.size[1] - 1],
        radius,
        outline=(246, 220, 155, 90),
        width=3,
    )
    canvas = Image.alpha_composite(canvas, border)
    canvas.save(dst)
    print("shot:", dst)


# ---------------------------------------------------------------- pptx utils
def set_font(run, name: str, size: float, color: RGBColor, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def spacing(run, pts: float):
    run._r.get_or_add_rPr().set("spc", str(int(pts * 100)))


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.25, space_after=0):
    """lines: list of dicts {text,size,color,font,bold,spacing,italic,space_before}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ln.get("line_spacing", line_spacing)
        p.space_after = Pt(ln.get("space_after", space_after))
        if "space_before" in ln:
            p.space_before = Pt(ln["space_before"])
        run = p.add_run()
        run.text = ln["text"]
        set_font(run, ln.get("font", SANS), ln["size"], ln["color"],
                 ln.get("bold", False), ln.get("italic", False))
        if "spacing" in ln:
            spacing(run, ln["spacing"])
    return box


def add_panel(slide, x, y, w, h, fill=PANEL, alpha=42, line=LINE, line_alpha=70, radius=0.16):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = radius
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    fill_el = shape.fill._xPr.find(qn("a:solidFill"))
    if fill_el is not None:
        clr = fill_el.find(qn("a:srgbClr"))
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))})
        clr.append(a)
    shape.line.color.rgb = line
    shape.line.width = Pt(0.75)
    ln = shape.line._get_or_add_ln()
    clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(line_alpha * 1000))}))
    return shape


def add_bg(slide, img_path: Path):
    pic = slide.shapes.add_picture(str(img_path), 0, 0, width=Inches(W), height=Inches(H))
    sp = pic._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def add_pic(slide, path: Path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def chip(slide, x, y, w, text, color=GOLD2, size=9.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.3)
    )
    shape.adjustments[0] = 0.5
    shape.shadow.inherit = False
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(0.75)
    tf = shape.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, SANS, size, color)
    spacing(r, 0.6)
    return shape


def emblem(slide, x, y, d):
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    ring.shadow.inherit = False
    ring.fill.background()
    ring.line.color.rgb = GOLD2
    ring.line.width = Pt(0.9)
    ln = ring.line._get_or_add_ln()
    c = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    c.append(c.makeelement(qn("a:alpha"), {"val": "45000"}))

    moon = slide.shapes.add_shape(
        MSO_SHAPE.MOON, Inches(x + d * 0.24), Inches(y + d * 0.2), Inches(d * 0.36), Inches(d * 0.42)
    )
    moon.rotation = 200
    moon.shadow.inherit = False
    moon.fill.solid()
    moon.fill.fore_color.rgb = GOLD
    moon.line.fill.background()

    star = slide.shapes.add_shape(
        MSO_SHAPE.STAR_4_POINT, Inches(x + d * 0.5), Inches(y + d * 0.44), Inches(d * 0.3), Inches(d * 0.3)
    )
    star.shadow.inherit = False
    star.fill.solid()
    star.fill.fore_color.rgb = FIRE
    star.line.fill.background()


def eyebrow(slide, text, y=0.62):
    add_text(slide, LEFT, y, 7, 0.3, [{
        "text": text, "size": 10, "color": GOLD3, "spacing": 2.2,
    }])


def page_no(slide, idx, total=10):
    add_text(slide, W - 1.55, H - 0.62, 0.9, 0.3, [{
        "text": f"{idx:02d} / {total:02d}", "size": 9, "color": MUTED, "spacing": 1.2,
    }], align=PP_ALIGN.RIGHT)


def title(slide, text, y=1.02, size=34, w=10.6):
    add_text(slide, LEFT, y, w, 0.8, [{
        "text": text, "size": size, "color": GOLD, "font": SERIF, "bold": True, "spacing": 1.0,
    }])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- deck
def build():
    make_background(str(DECK / "bg-cover.png"), cover=True)
    make_background(str(DECK / "bg-content.png"), cover=False)
    shot_dir = DECK / "shots"
    shot_dir.mkdir(exist_ok=True)
    rounded_shot(ASSETS / "spark-overview.png", shot_dir / "streak.png")
    rounded_shot(ASSETS / "focus-sound-v060.png", shot_dir / "focus.png")
    rounded_shot(ASSETS / "focus-policy-v060.png", shot_dir / "policy.png")
    rounded_shot(ASSETS / "spark-ai.png", shot_dir / "ai.png")

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    bg_cover = DECK / "bg-cover.png"
    bg_content = DECK / "bg-content.png"

    # ---- 1 cover
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_cover)
    emblem(s, W / 2 - 0.43, 1.16, 0.86)
    add_text(s, 0, 2.2, W, 0.3, [{
        "text": "ESTY  ·  SERIES", "size": 11, "color": GOLD2, "spacing": 4.2,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 2.62, W, 1.1, [{
        "text": "星火三部曲", "size": 58, "color": GOLD, "font": SERIF, "bold": True, "spacing": 3.0,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 3.82, W, 0.4, [{
        "text": "东风夜放花千树，更吹落，星如雨。", "size": 14, "color": GOLD, "font": SERIF, "spacing": 1.4,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 4.42, W, 0.4, [{
        "text": "从一份牵挂，到一段专注，再到更远的远方", "size": 12.5, "color": MUTED, "spacing": 0.6,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 5.95, W, 0.35, [{
        "text": "星 星 之 火  ·  可 以 燎 原", "size": 11, "color": FIRE2, "font": SERIF, "spacing": 2.0,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 6.62, W, 0.3, [{
        "text": "STARFIRE TRILOGY  ·  2026", "size": 8.5, "color": MUTED, "spacing": 2.0,
    }], align=PP_ALIGN.CENTER)
    notes(s, "开场：三部曲是写给未来的三封信——牵挂、专注、以及尚未命名的远方。")

    # ---- 2 meaning
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "01  /  何 为 星 火")
    title(s, "微小的开始，燎原的复利")
    cards = [
        ("壹", "星火", "微小而珍贵的开始", "一段火花、一次专注、一个习惯——任何宏大的改变，都始于不被看见的星火。"),
        ("贰", "续燃", "坚持是最大的杠杆", "火种的价值在于延续。日复一日地续上火花、守住专注，时间自会给出答案。"),
        ("叁", "燎原", "由点及面的成长", "一部作品是一颗火种，三部曲连成一片原野。记录技术成长与生活热爱。"),
    ]
    cw, gap = 3.6, 0.38
    for i, (num, name, head, body) in enumerate(cards):
        x = LEFT + i * (cw + gap)
        add_panel(s, x, 2.28, cw, 2.95)
        add_text(s, x + 0.36, 2.58, 1.2, 0.5, [{
            "text": num, "size": 26, "color": GOLD2, "font": SERIF, "bold": True,
        }])
        add_text(s, x + 0.36, 3.18, cw - 0.72, 0.35, [{
            "text": name, "size": 11, "color": GOLD3, "spacing": 3.0,
        }])
        add_text(s, x + 0.36, 3.62, cw - 0.72, 0.5, [{
            "text": head, "size": 16, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, x + 0.36, 4.22, cw - 0.72, 1.6, [{
            "text": body, "size": 10.5, "color": MUTED, "line_spacing": 1.55,
        }])
    add_text(s, LEFT, 5.72, 9, 0.35, [{
        "text": "星火不是宏大的宣言，而是一次次可以被重复的小动作。", "size": 10.5, "color": GOLD3, "font": SERIF, "italic": True,
    }])
    page_no(s, 2)
    notes(s, "把「星火」拆成三个层次：开始、延续、扩散。")

    # ---- 3 trilogy overview
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "02  /  三 部 曲")
    title(s, "一条关于「时间」的主线")
    rows = [
        ("01", "星火续火花", "STARFIRE · STREAK KEEPER", "已上线", GOLD2,
         "云端守护抖音火花，把每天一句问候交给程序。"),
        ("02", "星火专注 Pro", "STARFIRE · FOCUS PRO", "v0.6.0", GOLD2,
         "三种启动协议、判例簿、静场与国策树，把专注变成可训练的系统。"),
        ("03", "第三部", "TO BE CONTINUED", "构思中", MUTED,
         "火已燎原，方向未定；下一簇火，值得被认真点燃。"),
    ]
    for i, (num, name, en, status, scol, desc) in enumerate(rows):
        y = 2.18 + i * 1.52
        add_panel(s, LEFT, y, CONTENT_W, 1.26, alpha=34)
        add_text(s, LEFT + 0.35, y + 0.28, 0.8, 0.6, [{
            "text": num, "size": 24, "color": GOLD2, "font": SERIF, "bold": True,
        }])
        add_text(s, LEFT + 1.25, y + 0.22, 3.6, 0.36, [{
            "text": name, "size": 17, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, LEFT + 1.25, y + 0.68, 3.6, 0.3, [{
            "text": en, "size": 8.5, "color": GOLD3, "spacing": 1.4,
        }])
        add_text(s, LEFT + 5.1, y + 0.3, 5.2, 0.7, [{
            "text": desc, "size": 11, "color": MUTED, "line_spacing": 1.5,
        }], anchor=MSO_ANCHOR.MIDDLE)
        chip(s, LEFT + CONTENT_W - 1.35, y + 0.48, 1.0, status, scol, 9)
    page_no(s, 3)
    notes(s, "三个作品分别是：守护牵挂、训练专注、尚未命名。")

    # ---- 4 project 1
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "第 一 部  /  星 火 续 火 花")
    title(s, "让每一段牵挂，不断线", size=32, w=6.4)
    lines = [
        ("云端 24 小时待命", "关机、断网都不影响；多账号独立定时，并发自动排队。"),
        ("火花四态精准识别", "已续 / 待续 / 重燃 / 未点亮，一眼看清今天谁还差一句。"),
        ("日志可视化", "按级别着色、可搜索、只看错误，排查不再大海捞针。"),
    ]
    for i, (head, body) in enumerate(lines):
        y = 2.35 + i * 1.32
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(LEFT), Inches(y + 0.07), Inches(0.11), Inches(0.11))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = GOLD2
        dot.line.fill.background()
        add_text(s, LEFT + 0.32, y - 0.04, 5.5, 0.35, [{
            "text": head, "size": 14, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, LEFT + 0.32, y + 0.36, 5.5, 0.7, [{
            "text": body, "size": 10.5, "color": MUTED, "line_spacing": 1.5,
        }])
    add_pic(s, shot_dir / "streak.png", 7.05, 1.5, 5.55, 3.66)
    add_text(s, 7.05, 5.62, 5.55, 0.3, [{
        "text": "火花守护榜与当日状态", "size": 9, "color": MUTED,
    }], align=PP_ALIGN.CENTER)
    chip(s, 7.05, 6.25, 1.55, "V1.0 → V5.0", GOLD3, 9)
    chip(s, 8.8, 6.25, 1.35, "云端挂机", GOLD3, 9)
    page_no(s, 4)
    notes(s, "第一部解决的是「每天记得发消息」这件小事。")

    # ---- 5 AI engine
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "第 一 部  /  A I  引 擎")
    title(s, "让问候有温度", size=32)
    add_text(s, LEFT, 1.72, 10.4, 0.4, [{
        "text": "续火花的智能中枢——不再只是冷冰冰的一个表情。", "size": 11.5, "color": MUTED,
    }])
    ai = [
        ("01", "日常问候", "每天现写一句、不重样的问候；失败或超时自动回退本地语库，绝不漏发。"),
        ("02", "新闻早报", "早安 + 好友城市真实天气 + 当天新闻，按你的语气润色成专属早报。"),
        ("03", "智能接话", "读对方最近几条消息顺着回，像朋友一样自然；红灯话题引导到微信。"),
    ]
    cw2, gap2 = 3.6, 0.38
    for i, (num, head, body) in enumerate(ai):
        x = LEFT + i * (cw2 + gap2)
        add_panel(s, x, 2.45, cw2, 2.25)
        add_text(s, x + 0.34, 2.72, 1.0, 0.3, [{
            "text": num, "size": 9.5, "color": GOLD3, "spacing": 1.2,
        }])
        add_text(s, x + 0.34, 3.02, cw2 - 0.68, 0.4, [{
            "text": head, "size": 15, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, x + 0.34, 3.5, cw2 - 0.68, 1.2, [{
            "text": body, "size": 10, "color": MUTED, "line_spacing": 1.55,
        }])
    add_text(s, LEFT, 5.28, 10.9, 0.7, [{
        "text": "隐私与安全：使用你自己的接口 Key，Key 只存在你的服务器（不进备份、不进日志）；"
                "接话模式对方消息只当次临时读取，不落盘。",
        "size": 9.5, "color": MUTED, "line_spacing": 1.5,
    }])
    page_no(s, 5)
    notes(s, "AI 是第二层：让机械的续火花变得像人写的。")

    # ---- 6 project 2
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "第 二 部  /  星 火 专 注  P r o", )
    title(s, "把「开始」拆成三种协议", size=32, w=6.3)
    lines = [
        ("协议点火", "立即开始、5 分钟侦查或预约入座，先跨过启动阻力。"),
        ("计时 · 判例", "时间戳防漂移；早停时断链，或立为永久判例。"),
        ("国策 · 稳态", "每天最多点亮一条规则，用内化进度改写默认状态。"),
    ]
    for i, (head, body) in enumerate(lines):
        y = 2.35 + i * 1.32
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(LEFT), Inches(y + 0.07), Inches(0.11), Inches(0.11))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = GOLD2
        dot.line.fill.background()
        add_text(s, LEFT + 0.32, y - 0.04, 5.4, 0.35, [{
            "text": head, "size": 14, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, LEFT + 0.32, y + 0.36, 5.4, 0.7, [{
            "text": body, "size": 10.5, "color": MUTED, "line_spacing": 1.5,
        }])
    add_pic(s, shot_dir / "focus.png", 7.0, 1.5, 5.6, 3.69)
    add_text(s, 7.0, 5.66, 5.6, 0.3, [{
        "text": "v0.6.0 点火页 · 协议与本地静场", "size": 9, "color": MUTED,
    }], align=PP_ALIGN.CENTER)
    chip(s, 7.0, 6.25, 1.3, "可安装 PWA", GOLD3, 9)
    chip(s, 8.46, 6.25, 1.6, "数据只留本机", GOLD3, 9)
    chip(s, 10.22, 6.25, 1.28, "v0.6.0", GOLD3, 9)
    page_no(s, 6)
    notes(s, "第二部把「专注」从意志力问题，变成设计问题。")

    # ---- 7 loop
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "第 二 部  /  核 心 闭 环")
    title(s, "从一小时，走向长期稳态", size=32)
    steps = [
        ("壹", "协议点火", "把决定变成动作：选一种协议，先坐下。"),
        ("贰", "计时 · 判例", "早停不留退路：断链，或写成永久判例。"),
        ("叁", "国策 · 稳态", "每天点亮一条规则，慢慢改写默认状态。"),
    ]
    sw, sgap = 3.35, 0.7
    for i, (num, head, body) in enumerate(steps):
        x = LEFT + i * (sw + sgap)
        add_panel(s, x, 2.42, sw, 2.12, alpha=30)
        add_text(s, x + 0.3, 2.62, 1.0, 0.35, [{
            "text": num, "size": 20, "color": GOLD2, "font": SERIF, "bold": True,
        }])
        add_text(s, x + 0.3, 3.08, sw - 0.6, 0.4, [{
            "text": head, "size": 15, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, x + 0.3, 3.56, sw - 0.5, 1.1, [{
            "text": body, "size": 9.5, "color": MUTED, "line_spacing": 1.55,
        }])
        if i < 2:
            add_text(s, x + sw + 0.12, 3.38, 0.5, 0.4, [{
                "text": "→", "size": 18, "color": GOLD3,
            }], align=PP_ALIGN.CENTER)
    add_text(s, LEFT, 5.05, 9.6, 0.5, [{
        "text": "先让身体坐下，再等状态出现。", "size": 15, "color": GOLD, "font": SERIF, "italic": True,
    }])
    add_text(s, LEFT, 5.62, 9.6, 0.4, [{
        "text": "启动协议与判例机制受 edmond《如何提高自制力？》中 CTDP / RSIP 方法启发。",
        "size": 9, "color": MUTED,
    }])
    page_no(s, 7)
    notes(s, "闭环：点火 → 计时与判例 → 国策稳态。")

    # ---- 8 timeline
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_content)
    eyebrow(s, "0 4  /  星 火 年 表")
    title(s, "两个作品，一条向上的线", size=32)
    cols = [
        ("第一部 · 星火续火花", "2026 · 秋", GOLD2, [
            ("V1.0", "诞生：只有续火花"),
            ("V2.0", "语料库与节日应景"),
            ("V3.0", "一键上云 + AI 问候"),
            ("V4.0", "四态识别与守夜人告警"),
            ("V5.0", "最终版：日志与移动端"),
        ]),
        ("第二部 · 星火专注 Pro", "2026 · 秋", FIRE2, [
            ("v0.1", "立项：月夜星火"),
            ("v0.3", "星图 · 统计 · 备份"),
            ("v0.4", "协议点火 · 判例簿"),
            ("v0.5", "静场 · 品牌统一"),
            ("v0.6", "国策树 · 长期稳态"),
        ]),
    ]
    for ci, (head, when, accent, items) in enumerate(cols):
        x = LEFT + ci * 5.75
        add_text(s, x, 2.05, 4.6, 0.35, [{
            "text": head, "size": 14, "color": INK, "font": SERIF, "bold": True,
        }])
        add_text(s, x, 2.42, 4.6, 0.3, [{
            "text": when, "size": 9, "color": MUTED, "spacing": 1.2,
        }])
        for j, (ver, desc) in enumerate(items):
            y = 2.95 + j * 0.78
            dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.12), Inches(0.1), Inches(0.1))
            dot.shadow.inherit = False
            dot.fill.solid()
            dot.fill.fore_color.rgb = accent
            dot.line.fill.background()
            add_text(s, x + 0.26, y, 0.95, 0.32, [{
                "text": ver, "size": 9.5, "color": accent, "spacing": 0.8,
            }])
            add_text(s, x + 1.28, y - 0.01, 3.35, 0.34, [{
                "text": desc, "size": 10.5, "color": MUTED,
            }])
    page_no(s, 8)
    notes(s, "版本不是炫耀，是每一轮把问题修好的记录。")

    # ---- 9 epilogue
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_cover)
    emblem(s, W / 2 - 0.31, 1.12, 0.62)
    add_text(s, 0, 2.02, W, 0.3, [{
        "text": "第 三 部  ·  构 思 中", "size": 10, "color": GOLD3, "spacing": 3.0,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 1.6, 2.6, W - 3.2, 1.9, [
        {
            "text": "当技术把创造的火种递给普通人，", "size": 27, "color": GOLD,
            "font": SERIF, "bold": True, "line_spacing": 1.45,
        },
        {
            "text": "我们究竟能用它点亮什么？", "size": 27, "color": GOLD,
            "font": SERIF, "bold": True, "line_spacing": 1.45,
        },
    ], align=PP_ALIGN.CENTER)
    add_text(s, 1.6, 4.72, W - 3.2, 0.4, [{
        "text": "第三部写什么，我还没想好。但火种已经埋下。", "size": 11.5, "color": MUTED,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 6.25, W, 0.35, [{
        "text": "—— Esty，写于星火燎原的秋夜", "size": 10, "color": GOLD3, "font": SERIF,
    }], align=PP_ALIGN.CENTER)
    page_no(s, 9)
    notes(s, "结尾留一个开放的问题，而不是一个总结。")

    # ---- 10 back
    s = prs.slides.add_slide(blank)
    add_bg(s, bg_cover)
    add_text(s, 0, 2.42, W, 0.9, [{
        "text": "星星之火，可以燎原。", "size": 40, "color": GOLD, "font": SERIF, "bold": True, "spacing": 2.0,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 3.62, W, 0.35, [{
        "text": "愿你不忘来时的星火，也敢追更远的原野。", "size": 12, "color": MUTED,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 4.82, W, 0.35, [{
        "text": "beeasy-l.github.io/starspark-trilogy", "size": 11, "color": GOLD2, "font": "Consolas",
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 5.25, W, 0.35, [{
        "text": "第一部 星火续火花   ·   第二部 星火专注 Pro   ·   第三部 待定", "size": 10, "color": MUTED,
    }], align=PP_ALIGN.CENTER)
    add_text(s, 0, 6.55, W, 0.3, [{
        "text": "STARFIRE SERIES  ·  © 2026 ESTY", "size": 8.5, "color": GOLD3, "spacing": 2.0,
    }], align=PP_ALIGN.CENTER)
    notes(s, "收尾：站点地址与三部曲定位。")

    prs.core_properties.title = "星火三部曲 · 项目介绍"
    prs.core_properties.author = "Esty"
    prs.core_properties.subject = "星火三部曲系列项目介绍"
    prs.core_properties.comments = "由 DSH 的 pptx 技能生成"
    prs.save(str(OUT))
    print("saved:", OUT, "slides:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    build()
